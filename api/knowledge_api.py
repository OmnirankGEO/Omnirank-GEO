"""
统一知识库 API 路由
提供角色/客户知识库的上传、列表、删除功能
"""

from fastapi import APIRouter, HTTPException, UploadFile, File, Form, Request
from pydantic import BaseModel
from pathlib import Path
from typing import Any, Dict, Optional
from tools.unified_knowledge import get_unified_rag
import asyncio
import logging
import time
from types import SimpleNamespace

_logger = logging.getLogger("GEO-Knowledge-API")

router = APIRouter(prefix="/api/knowledge", tags=["统一知识库"])
BASE_DIR = Path(__file__).resolve().parent.parent
VECTORDB_ROOT = BASE_DIR / "data" / "vectordb"
CLIENT_KNOWLEDGE_ROOT = BASE_DIR / "data" / "knowledge" / "clients"
CLIENT_UPLOAD_FAST_PIPELINE_TIMEOUT_SECONDS = 8.0

# [GEO-R1-CAN-055] 上传体积/解析上限 · 防 zip-bomb 式解压放大拖垮单 worker(WORKERS=1)
KB_UPLOAD_MAX_BYTES = 20 * 1024 * 1024  # 原始上传字节上限 20MB
KB_EXTRACTED_TEXT_MAX_CHARS = 2_000_000  # 提取文本字符上限,防解压/解析放大


def _require_admin(request: Request) -> None:
    """[GEO-R1-CAN-054/057/064/126] 管理员闸 · 对齐 /public/upload 既有校验,fail-closed 防普通写作用户污染/枚举全局语料。"""
    user = getattr(request.state, "user", None)
    if not user or not user.get("is_admin"):
        raise HTTPException(status_code=403, detail="需要管理员权限")


class DocumentUpload(BaseModel):
    """文档上传请求"""
    kb_type: str  # "client" | "role"
    kb_id: Optional[str] = None  # brand_id 或 role_id
    role_type: Optional[str] = None  # "advisor" | "employee"
    filename: str
    content: str


class DocumentDelete(BaseModel):
    """文档删除请求"""
    kb_type: str  # "client" | "role"
    kb_id: Optional[str] = None
    role_type: Optional[str] = None
    filename: str


class KnowledgeSearch(BaseModel):
    """客户知识库语义检索请求"""
    query: str
    brand_id: Optional[int] = None
    top_k: int = 5
    use_rerank: bool = True


def _safe_filename(filename: str) -> str:
    """只保留上传文件名，避免写入客户知识库目录外。"""
    safe = Path(filename or "未命名资料.txt").name.strip()
    return safe or "未命名资料.txt"


def _safe_relative_path(filename: str) -> str:
    """允许删除客户知识库里的相对路径文件，但禁止目录逃逸。"""
    raw = (filename or "未命名资料.txt").replace("\\", "/").strip().lstrip("/")
    parts = [part for part in raw.split("/") if part and part != "."]
    if any(part == ".." for part in parts):
        raise HTTPException(status_code=400, detail="非法文件路径")
    safe_parts = [_safe_filename(part) for part in parts]
    return "/".join(safe_parts) or _safe_filename(filename)


def _client_vector_counts(brand_id: int) -> dict[str, int]:
    """读取客户 LanceDB 向量表，按文件名统计 chunk 数。"""
    try:
        import lancedb

        if not VECTORDB_ROOT.exists():
            return {}

        db = lancedb.connect(str(VECTORDB_ROOT))
        table_name = f"knowledge_client_{brand_id}"
        try:
            table = db.open_table(table_name)
        except Exception:
            return {}

        df = table.to_pandas()
        if df.empty or "filename" not in df.columns:
            return {}

        return {str(k): int(v) for k, v in df.groupby("filename").size().to_dict().items()}
    except Exception as e:
        _logger.warning(f"[client-kb] 读取向量统计失败 brand_id={brand_id}: {e}")
        return {}


def _enrich_client_documents(brand_id: int, documents: list[dict]) -> list[dict]:
    counts = _client_vector_counts(brand_id)
    enriched = []
    for doc in documents:
        filename = doc.get("filename", "")
        vector_chunks = counts.get(filename, 0)
        enriched.append({
            **doc,
            "vector_chunks": vector_chunks,
            "vectorized": vector_chunks > 0,
        })
    return enriched


def _pipeline_response(result: Any) -> Dict[str, Any]:
    return {
        "processed": bool(getattr(result, "success", False)),
        "chunk_count": int(getattr(result, "chunk_count", 0) or 0),
        "knowledge_points": int(getattr(result, "knowledge_point_count", 0) or 0),
        "keywords": list(getattr(result, "keywords", []) or []),
        "processing_time_ms": int(getattr(result, "processing_time_ms", 0) or 0),
    }


async def _process_client_document(brand_id: int, filename: str, content: str, skip_clean: bool = False):
    """客户资料必须进入 KnowledgePipeline，确保写入 knowledge_client_{brand_id} 向量表。"""
    try:
        from services.knowledge_pipeline import get_knowledge_pipeline
        pipeline = get_knowledge_pipeline()
        result = await pipeline.process_document(
            content=content,
            filename=_safe_filename(filename),
            brand_id=brand_id,
            skip_clean=skip_clean,
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"向量入库失败: {e}")

    if not result.success:
        raise HTTPException(status_code=500, detail=result.error or "向量入库失败")

    return result


async def _background_client_document_clean_and_sync(
    brand_id: int,
    filename: str,
    content: str,
) -> None:
    """后台补做慢清洗与客户资料同步；失败隔离，绝不影响上传响应。"""
    fallback_context = content
    try:
        result = await _process_client_document(brand_id, filename, content, skip_clean=False)
        if bool(getattr(result, "success", False)):
            fallback_context = ""
            _logger.info(
                f"[client-kb] 后台清洗/重建索引完成: brand_id={brand_id}, "
                f"filename={filename}, chunks={getattr(result, 'chunk_count', 0)}"
            )
    except Exception as e:
        _logger.warning(
            f"[client-kb] 后台清洗/重建索引失败(已忽略): brand_id={brand_id}, "
            f"filename={filename}, err={e}"
        )

    try:
        await _auto_trigger_material_clean(brand_id, fallback_context)
    except Exception as e:
        _logger.warning(
            f"[client-kb] 后台客户资料整理失败(已忽略): brand_id={brand_id}, "
            f"filename={filename}, err={e}"
        )


def _save_client_document_fallback(brand_id: int, filename: str, content: str, error: str = ""):
    """保存客户原文兜底：AI 整理/向量化失败不能让上传本身失败。"""
    rag = get_unified_rag()
    result = rag.add_document(
        kb_type="client",
        kb_id=str(brand_id),
        filename=_safe_filename(filename),
        content=content,
    )
    if not result.get("success"):
        raise HTTPException(status_code=500, detail=result.get("error", "资料保存失败"))

    return SimpleNamespace(
        success=True,
        original_path=result.get("path", ""),
        cleaned_path="",
        chunk_count=0,
        knowledge_point_count=0,
        keywords=[],
        processing_time_ms=0,
        error=error or "AI 整理暂未完成",
        vectorized=False,
    )


# ==================== 上传后自动整理(知识库同步根治 · P4) ====================
# 客户上传/向量化资料后，自动触发 M3 的 AI 整理逻辑，把新资料梳理进
# client_materials + client_profiles，让「诊断」「报价」读到的结构化资料能随上传更新。
# 复用 api/m3_material_confirm_api.py 的核心内部函数（_get_brand_info /
# _get_profile_for_brand / _get_materials_for_brand / _retrieve_client_knowledge_context /
# _ai_clean_materials / _persist_cleaned_materials），绝不走 HTTP 自调。

# 进程内防抖：同一 brand 在 _AUTO_CLEAN_DEBOUNCE_SECONDS 内只触发一次自动整理。
# clean 端点本身有限流(60s/3 次)；这里再加一层防抖，避免一次多文件上传连发多次整理。
_AUTO_CLEAN_DEBOUNCE_SECONDS = 90
_auto_clean_last_ts: Dict[int, float] = {}


async def _auto_trigger_material_clean(brand_id: int, fallback_context: str = "") -> None:
    """上传客户资料成功后台触发 AI 整理(失败全隔离 · 绝不影响上传响应)。

    复用 m3_material_confirm_api 的核心内部函数，不走 HTTP 自调：
      - _get_brand_info / _get_profile_for_brand / _get_materials_for_brand 取上下文
      - _retrieve_client_knowledge_context 从客户向量库检索刚上传的资料片段
      - _ai_clean_materials 调 LLM 梳理成结构化 JSON
      - _persist_cleaned_materials 写回 client_materials(+ client_profiles)
    """
    try:
        # 防抖:同一 brand 短时间内多次上传只整理一次
        now = time.time()
        last = _auto_clean_last_ts.get(brand_id, 0.0)
        if now - last < _AUTO_CLEAN_DEBOUNCE_SECONDS:
            _logger.info(
                f"[client-kb] 自动整理跳过(防抖): brand_id={brand_id} · "
                f"距上次仅 {int(now - last)}s(<{_AUTO_CLEAN_DEBOUNCE_SECONDS}s)"
            )
            return
        _auto_clean_last_ts[brand_id] = now
        # 简单内存治理，避免长期累积
        if len(_auto_clean_last_ts) > 5000:
            try:
                _auto_clean_last_ts.pop(next(iter(_auto_clean_last_ts)), None)
            except Exception:
                pass

        from api.m3_material_confirm_api import (
            _get_brand_info,
            _get_profile_for_brand,
            _get_materials_for_brand,
            _retrieve_client_knowledge_context,
            _ai_clean_materials,
            _persist_cleaned_materials,
        )

        brand_info = _get_brand_info(brand_id)
        profile = _get_profile_for_brand(brand_id)
        materials = _get_materials_for_brand(brand_id)

        # 从客户知识库(刚上传并向量化的资料)检索写作相关片段作为整理输入
        kb_context = await _retrieve_client_knowledge_context(brand_id, brand_info)
        if not kb_context.strip() and fallback_context.strip():
            kb_context = fallback_context.strip()
        if not kb_context.strip():
            _logger.info(f"[client-kb] 自动整理跳过: brand_id={brand_id} 知识库检索为空")
            return

        cleaned = await _ai_clean_materials(brand_info, profile, materials, kb_context, "")
        if cleaned.get("_llm_error"):
            _logger.warning(
                f"[client-kb] 自动整理 LLM 失败 brand_id={brand_id}: "
                f"{cleaned.get('warnings')}"
            )
            return

        # diagnosis_id 传 None → _persist 内部自动找最新诊断兜底 / 退 client_profiles 路径
        _persist_cleaned_materials(brand_id, None, cleaned)
        _logger.info(f"[client-kb] 上传后自动整理完成: brand_id={brand_id}")
    except Exception as e:
        # 失败完全隔离：只 log warning，绝不向上抛(后台任务不影响上传成功响应)
        _logger.warning(f"[client-kb] 上传后自动整理失败(已忽略) brand_id={brand_id}: {e}")


def _client_document_response(result: Any, filename: str, content: str) -> Dict[str, Any]:
    vectorized = bool(getattr(result, "vectorized", True))
    pipeline_error = getattr(result, "error", None)
    return {
        "status": "success",
        "success": True,
        "message": "客户资料已上传、清洗并向量化" if vectorized else "客户资料已保存，AI 整理稍后可刷新查看",
        "path": getattr(result, "original_path", ""),
        "size": len(content.encode("utf-8")),
        "filename": _safe_filename(filename),
        "text_length": len(content),
        "vectorized": vectorized,
        "chunk_count": int(getattr(result, "chunk_count", 0) or 0),
        "processing_time_ms": int(getattr(result, "processing_time_ms", 0) or 0),
        "pipeline": {
            **_pipeline_response(result),
            "processed": vectorized and bool(getattr(result, "success", False)),
            "error": pipeline_error if not vectorized else None,
        },
    }


async def _extract_upload_text(file: UploadFile, filename: str) -> str:
    """从上传文件读文本(兼容旧调用)· 内部走 _extract_text_from_bytes"""
    content = await file.read()
    return _extract_text_from_bytes(content, filename)


def _extract_text_from_bytes(content: bytes, filename: str) -> str:
    """从已读 bytes 提文本(供 upload_file_document 一次读、文本+内嵌图共用同一份 bytes)"""
    ext = Path(filename).suffix.lower()
    # 旧格式 Word/PPT(.doc/.ppt 是 2007 年前的二进制格式)系统读不了,给人话引导别让用户看英文报错
    legacy_office = {".doc": "Word", ".ppt": "PPT"}
    if ext in legacy_office:
        raise HTTPException(
            status_code=400,
            detail=f"这是较老的 {legacy_office[ext]} 旧格式,系统暂时认不了。请在 {legacy_office[ext]} 里点【另存为】,选新版格式(.{'docx' if ext == '.doc' else 'pptx'})再上传就行啦~",
        )
    allowed = {".pdf", ".docx", ".txt", ".md", ".pptx", ".csv", ".json"}
    if ext not in allowed:
        raise HTTPException(status_code=400, detail=f"暂不支持这种文件({ext or '未知'}),换 Word(.docx)、PDF 或文本文件试试吧~")

    try:
        if ext in {".txt", ".md", ".csv", ".json"}:
            return content.decode("utf-8", errors="ignore")
        if ext == ".pdf":
            from services import pdf_reader
            return "\n\n".join(pdf_reader.page_texts(content))
        if ext == ".docx":
            import docx
            import io
            doc = docx.Document(io.BytesIO(content))
            return "\n".join(para.text for para in doc.paragraphs)
        if ext == ".pptx":
            from pptx import Presentation
            import io
            prs = Presentation(io.BytesIO(content))
            parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    text = getattr(shape, "text", "")
                    if text:
                        parts.append(text)
            return "\n".join(parts)
    except HTTPException:
        raise
    except Exception as e:
        # 真实错误只记后端日志,不把工程细节(如 docx PackageNotFoundError)暴露给用户
        _logger.warning(f"[KB] 文件读取失败 {filename}: {e}")
        raise HTTPException(status_code=400, detail="这个文件没能读出内容,可能是格式较特殊或文件损坏了。换一个文件,或另存为 .docx / .pdf 再试试吧~")

    return ""


# [任务2 2026-06-04] 文档内嵌图 → 图片素材库(复用 image_asset 存图 + vision 打标 + 落表)
KB_EMBEDDED_IMAGE_MAX = 10          # 单文档最多入库 N 张内嵌图 · 防爆 vision 成本(对齐 social PDF_MAX_VISION_IMAGES 理念)
KB_EMBEDDED_IMAGE_MIN_BYTES = 4096  # < 4KB 当 icon/装饰跳过


def _extract_embedded_images(raw: bytes, filename: str, max_images: int = KB_EMBEDDED_IMAGE_MAX) -> list:
    """从 PDF / docx 提**内嵌二进制图**(外链 ![](url) 不管 · 过滤 icon · 限张数)· 返回 [(img_bytes, img_filename)]"""
    ext = Path(filename).suffix.lower()
    out: list = []
    try:
        if ext == ".pdf":
            from services import pdf_reader
            for img in pdf_reader.embedded_images(raw, KB_EMBEDDED_IMAGE_MIN_BYTES, max_images):
                out.append((img.data, f"{filename}_p{img.page_index + 1}_img{len(out) + 1}.{img.ext}"))
        elif ext == ".docx":
            import zipfile
            import io
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                for name in z.namelist():
                    if len(out) >= max_images:
                        break
                    if name.startswith("word/media/"):
                        b = z.read(name)
                        if b and len(b) >= KB_EMBEDDED_IMAGE_MIN_BYTES:
                            e = name.rsplit(".", 1)[-1].lower() if "." in name else "png"
                            out.append((b, f"{filename}_img{len(out) + 1}.{e}"))
    except Exception as e:
        _logger.warning(f"[client-kb] 内嵌图提取失败(忽略) filename={filename}: {e}")
    return out


async def _persist_embedded_images_to_library(brand_id: int, images: list, uploaded_by: Optional[int] = None) -> int:
    """内嵌图复用图片素材库存图流程(存储 + vision 打标 + 落表)· 单图失败不阻断"""
    if not images:
        return 0
    try:
        from db.brand_image_assets_db import init_brand_image_assets_table, insert_image_asset
        from services.image_storage import save_brand_image
        from tools.vision.image_describe import describe_image_structured
    except Exception as e:
        _logger.warning(f"[client-kb] 内嵌图入库依赖加载失败(忽略) brand_id={brand_id}: {e}")
        return 0
    try:
        init_brand_image_assets_table()
    except Exception:
        pass
    n = 0
    for img_bytes, img_name in images:
        try:
            stored = save_brand_image(brand_id, img_bytes, img_name)
            desc = await describe_image_structured(img_bytes, img_name)
            insert_image_asset(
                brand_id=brand_id, file_name=img_name,
                storage_key=stored["storage_key"], public_url=stored["public_url"],
                thumbnail_key=stored.get("thumbnail_key"), safe_size_key=stored.get("safe_size_key"),
                image_type=desc.get("image_type"), title=desc.get("title"), caption=desc.get("caption"),
                alt_text=desc.get("alt_text"), suggested_placement=desc.get("suggested_placement"),
                usage_scenarios=desc.get("usage_scenarios"), vision_summary=desc.get("vision_summary"),
                ocr_text=desc.get("ocr_text"), tags=desc.get("tags"),
                publish_allowed=desc.get("publish_allowed", 0), rights_confirmed=0,
                risk_flags=desc.get("risk_flags"), uploaded_by=uploaded_by,
            )
            n += 1
        except Exception as e:
            _logger.warning(f"[client-kb] 内嵌图入库失败(忽略) brand_id={brand_id} img={img_name}: {e}")
    _logger.info(f"[client-kb] 文档内嵌图入素材库 brand_id={brand_id}: {n}/{len(images)} 张")
    return n


async def _extract_text_from_upload(file: UploadFile) -> tuple[str, str]:
    """把上传文件解析成文本，供通用上传接口使用。"""
    filename = _safe_filename(file.filename or "未命名资料.txt")
    text = (await _extract_upload_text(file, filename)).strip()
    if not text:
        raise HTTPException(status_code=400, detail="无法从文件中提取文本内容")
    return filename, text


def _parse_client_brand_id(kb_id: Optional[str]) -> int:
    try:
        brand_id = int(kb_id or "")
    except Exception:
        raise HTTPException(status_code=400, detail="client 知识库需要有效 kb_id")
    if brand_id <= 0:
        raise HTTPException(status_code=400, detail="client 知识库需要有效 kb_id")
    return brand_id


# ==================== 角色知识库 API ====================

@router.get("/role/{role_type}/{role_id}", summary="获取角色知识库文档列表")
async def list_role_documents(role_type: str, role_id: str, http_request: Request):
    """
    获取顾问或员工的私有知识库文档列表

    - role_type: advisor | employee
    - role_id: 角色ID
    """
    # [GEO-R1-CAN-126] 角色(advisor/employee)知识库为平台全局命名空间,枚举需管理员
    _require_admin(http_request)
    rag = get_unified_rag()
    result = rag.list_documents(kb_type="role", kb_id=role_id, role_type=role_type)
    
    if not result.get("success"):
        raise HTTPException(status_code=500, detail=result.get("error", "获取失败"))
    
    return {
        "success": True,
        "role_type": role_type,
        "role_id": role_id,
        "documents": result.get("documents", []),
        "total": result.get("total", 0)
    }


@router.post("/role/{role_type}/{role_id}/upload", summary="上传文档到角色知识库")
async def upload_role_document(
    role_type: str,
    role_id: str,
    http_request: Request,
    file: UploadFile = File(...)
):
    """
    上传文档到顾问或员工的私有知识库（支持md/txt/json/pdf）

    完整处理流程：
    1. LLM清洗提取知识点
    2. 向量化（批次10，并发15）
    3. 存储到LanceDB
    """
    # [对抗审核订正 R1-CAN-054] 角色语料是全站 advisor/employee 共享检索源 → 仅 admin 可写。
    _require_admin(http_request)
    from services.knowledge_pipeline import get_knowledge_pipeline
    
    filename = file.filename or "unknown.txt"
    
    try:
        content = await file.read()
        
        # 根据文件类型解析内容
        if filename.lower().endswith('.pdf'):
            try:
                from services import pdf_reader
                content_str = "\n\n".join(pdf_reader.page_texts(content))
            except ImportError:
                raise HTTPException(status_code=500, detail="服务器未安装PDF解析库")
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"PDF解析失败: {e}")
        else:
            content_str = content.decode("utf-8")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件读取失败: {e}")
    
    # 使用完整Pipeline处理
    pipeline = get_knowledge_pipeline()
    result = await pipeline.process_role_document(
        content=content_str,
        filename=filename,
        role_type=role_type,
        role_id=role_id,
        skip_clean=False  # 走完整清洗流程
    )
    
    if not result.success:
        raise HTTPException(status_code=500, detail=result.error or "处理失败")
    
    return {
        "success": True,
        "message": f"文档已处理并存储到 {role_type}/{role_id}",
        "path": result.original_path,
        "chunk_count": result.chunk_count,
        "knowledge_points": result.knowledge_point_count,
        "processing_time_ms": result.processing_time_ms
    }



@router.delete("/role/{role_type}/{role_id}/{filename}", summary="删除角色知识库文档")
async def delete_role_document(role_type: str, role_id: str, filename: str, http_request: Request):
    """删除顾问或员工的私有知识库文档"""
    # [Deploy-CTO NO-GO finding 3 返工 2026-07-12] 角色语料(advisor/employee)是全站共享检索源,
    # 删除破坏性且跨租户 → 仅 admin(对齐 upload_role_document 的 _require_admin 写闸)。
    # 原实现无任何鉴权参数,任意登录用户可删他人角色知识库文档。
    _require_admin(http_request)
    rag = get_unified_rag()
    result = rag.delete_document(
        kb_type="role",
        kb_id=role_id,
        role_type=role_type,
        filename=_safe_filename(filename)
    )
    
    if not result.get("success"):
        raise HTTPException(status_code=500, detail=result.get("error", "删除失败"))
    
    return {"success": True, "message": "文档已删除"}


# ==================== 客户知识库 API ====================

@router.get("/client", summary="获取客户知识库文档列表(兼容 kb_id 查询)")
async def list_client_documents_by_query(kb_id: str, request: Request):
    """兼容旧前端：GET /api/knowledge/client?kb_id={brand_id}。"""
    if not kb_id:
        raise HTTPException(status_code=400, detail="缺少 kb_id")

    brand_id = _parse_client_brand_id(kb_id)
    from auth.brand_access import require_brand_access
    require_brand_access(request, brand_id)

    rag = get_unified_rag()
    result = rag.list_documents(kb_type="client", kb_id=str(brand_id))

    if not result.get("success"):
        raise HTTPException(status_code=500, detail=result.get("error", "获取失败"))

    return {
        "success": True,
        "brand_id": brand_id,
        "documents": _enrich_client_documents(brand_id, result.get("documents", [])),
        "total": result.get("total", 0)
    }


@router.get("/client/{brand_id}", summary="获取客户知识库文档列表")
async def list_client_documents(brand_id: int, request: Request):
    """获取客户知识库文档列表"""
    from auth.brand_access import require_brand_access
    require_brand_access(request, brand_id)

    rag = get_unified_rag()
    result = rag.list_documents(kb_type="client", kb_id=str(brand_id))
    
    if not result.get("success"):
        raise HTTPException(status_code=500, detail=result.get("error", "获取失败"))
    
    return {
        "success": True,
        "brand_id": brand_id,
        "documents": _enrich_client_documents(brand_id, result.get("documents", [])),
        "total": result.get("total", 0)
    }


@router.post("/upload", summary="保存文本到统一知识库")
async def upload_document(data: DocumentUpload, request: Request):
    """保存已经提取好的文本资料，支持客户/角色/公共知识库。"""
    filename = _safe_filename(data.filename)
    content = (data.content or "").strip()
    if not content:
        raise HTTPException(status_code=400, detail="资料内容为空")

    if data.kb_type == "client":
        brand_id = _parse_client_brand_id(data.kb_id)
        from auth.brand_access import require_brand_access
        require_brand_access(request, brand_id)
        result = await _process_client_document(brand_id, filename, content, skip_clean=False)
        # 上传成功后后台自动整理(知识库同步根治 · P4)· 失败隔离 · 不影响上传响应
        try:
            asyncio.create_task(_auto_trigger_material_clean(brand_id))
        except Exception as _bg_e:
            _logger.warning(f"[client-kb] 自动整理后台任务创建失败(已忽略) brand_id={brand_id}: {_bg_e}")
        return _client_document_response(result, filename, content)

    # [对抗审核订正 R1-CAN-054] 非 client(public/role)写全站共享语料 → 必须 admin。
    # 与 upload_file_document 的 public 分支闸对齐,堵住经本平行端点污染全站 AI 检索。
    _require_admin(request)
    rag = get_unified_rag()
    result = rag.add_document(
        kb_type=data.kb_type,
        kb_id=str(data.kb_id or ""),
        role_type=data.role_type,
        filename=filename,
        content=content,
    )

    if not result.get("success"):
        raise HTTPException(status_code=500, detail=result.get("error", "保存失败"))

    return {
        "success": True,
        "message": f"资料 {filename} 已保存",
        "path": result.get("path", ""),
        "size": result.get("size", 0),
        "filename": filename,
    }


@router.post("/upload-file", summary="上传文件到统一知识库")
async def upload_file_document(
    request: Request,
    file: UploadFile = File(...),
    kb_type: str = Form(...),
    kb_id: Optional[str] = Form(None),
    role_type: Optional[str] = Form(None),
):
    """上传 PDF/Word/TXT/Markdown/PPTX 到统一知识库，并以可检索文本保存。"""
    filename = _safe_filename(file.filename or "未命名资料.txt")
    raw_bytes = await file.read()  # [任务2] 一次读 · 文本提取 + 内嵌图提取共用同一份 bytes
    # [GEO-R1-CAN-055] 解析前先卡原始体积上限,防 zip-bomb 式解压放大在文本提取阶段拖垮单 worker
    if len(raw_bytes) > KB_UPLOAD_MAX_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"文件太大啦(超过 {KB_UPLOAD_MAX_BYTES // (1024 * 1024)}MB),请压缩或拆分后再上传~",
        )
    content = _extract_text_from_bytes(raw_bytes, filename).strip()
    if not content:
        raise HTTPException(status_code=400, detail="无法从文件中提取文本内容")
    # [GEO-R1-CAN-055] 再卡提取后文本上限,防解压/解析放大后的超大文本进入向量入库
    if len(content) > KB_EXTRACTED_TEXT_MAX_CHARS:
        content = content[:KB_EXTRACTED_TEXT_MAX_CHARS]

    if kb_type == "client":
        brand_id = _parse_client_brand_id(kb_id)
        from auth.brand_access import require_brand_access
        require_brand_access(request, brand_id)
        try:
            result = await asyncio.wait_for(
                _process_client_document(brand_id, filename, content, skip_clean=True),
                timeout=CLIENT_UPLOAD_FAST_PIPELINE_TIMEOUT_SECONDS,
            )
        except TimeoutError:
            _logger.warning(
                f"[client-kb] 快速向量入库超时, 已保留原文并转后台: "
                f"brand_id={brand_id}, filename={filename}"
            )
            result = _save_client_document_fallback(
                brand_id,
                filename,
                content,
                "快速向量入库超时，已转后台处理",
            )
        except HTTPException as exc:
            if exc.status_code < 500:
                raise
            _logger.warning(
                f"[client-kb] 向量/整理失败但保留原文: brand_id={brand_id}, "
                f"filename={filename}, err={exc.detail}"
            )
            result = _save_client_document_fallback(brand_id, filename, content, str(exc.detail))
        if bool(getattr(result, "vectorized", True)):
            _logger.info(
                f"[client-kb] 上传并向量入库: brand_id={brand_id}, filename={filename}, chunks={result.chunk_count}"
            )
        else:
            _logger.info(
                f"[client-kb] 上传已保存(向量暂未完成): brand_id={brand_id}, filename={filename}"
            )
        # [任务2 2026-06-04] 文档内嵌图自动入图片素材库(后台 · vision 打标 · 失败隔离不阻断上传)· 文字照旧向量化(上方)
        try:
            _embedded = _extract_embedded_images(raw_bytes, filename)
            if _embedded:
                _user = getattr(request.state, "user", None) or {}
                _uid = _user.get("user_id") if isinstance(_user, dict) else None
                asyncio.create_task(_persist_embedded_images_to_library(brand_id, _embedded, _uid))
        except Exception as _img_e:
            _logger.warning(f"[client-kb] 内嵌图入库调度失败(忽略) brand_id={brand_id}: {_img_e}")
        # 上传成功后后台补做慢清洗 + 自动整理(失败隔离 · 不影响上传响应)
        try:
            asyncio.create_task(_background_client_document_clean_and_sync(brand_id, filename, content))
        except Exception as _bg_e:
            _logger.warning(f"[client-kb] 自动整理后台任务创建失败(已忽略) brand_id={brand_id}: {_bg_e}")
        return _client_document_response(result, filename, content)

    # [GEO-R1-CAN-054] public/role 全局语料写入需管理员,对齐 /public/upload 的 is_admin 闸,
    # 防普通写作用户(agent/customer)经 kb_type=public 污染喂给全站 AI 检索的公共语料。
    _require_admin(request)

    rag = get_unified_rag()
    result = rag.add_document(
        kb_type=kb_type,
        kb_id=str(kb_id or ""),
        role_type=role_type,
        filename=filename,
        content=content,
    )

    if not result.get("success"):
        raise HTTPException(status_code=500, detail=result.get("error", "上传失败"))

    _logger.info(f"[unified-kb] 上传文档: type={kb_type}, kb_id={kb_id}, filename={filename}")
    return {
        "success": True,
        "message": f"文档 {filename} 已上传",
        "path": result.get("path", ""),
        "size": result.get("size", 0),
        "filename": filename,
        "text_length": len(content),
    }


@router.post("/client/{brand_id}/reindex", summary="重建客户知识库向量索引")
async def reindex_client_documents(brand_id: int, request: Request):
    """把客户知识库目录中的已有文件重新写入 LanceDB，修复历史未向量化文件。"""
    from auth.brand_access import require_brand_access
    require_brand_access(request, brand_id)

    rag = get_unified_rag()
    result = rag.list_documents(kb_type="client", kb_id=str(brand_id))
    if not result.get("success"):
        raise HTTPException(status_code=500, detail=result.get("error", "获取失败"))

    client_dir = CLIENT_KNOWLEDGE_ROOT / str(brand_id)
    indexed = 0
    failed: list[dict] = []

    for doc in result.get("documents", []):
        relative_path = doc.get("relative_path") or doc.get("filename")
        filename = _safe_filename(doc.get("filename") or relative_path)
        if not relative_path or str(relative_path).startswith("cleaned/"):
            continue

        path = client_dir / str(relative_path)
        if not path.exists() or not path.is_file():
            continue

        try:
            content = path.read_text(encoding="utf-8", errors="ignore").strip()
            if not content:
                failed.append({"filename": filename, "error": "文件内容为空"})
                continue
            await _process_client_document(brand_id, filename, content, skip_clean=True)
            indexed += 1
        except HTTPException as e:
            failed.append({"filename": filename, "error": str(e.detail)})
        except Exception as e:
            failed.append({"filename": filename, "error": str(e)})

    if indexed == 0 and failed:
        raise HTTPException(status_code=500, detail={"message": "重建索引失败", "failed": failed})

    return {
        "success": True,
        "brand_id": brand_id,
        "indexed": indexed,
        "failed": failed,
    }


@router.delete("/client/{brand_id}/{filename:path}", summary="删除客户知识库文档")
async def delete_client_document(brand_id: int, filename: str, request: Request):
    """删除客户知识库中的指定文档。"""
    from auth.brand_access import require_brand_access
    require_brand_access(request, brand_id)

    rag = get_unified_rag()
    safe_path = _safe_relative_path(filename)
    result = rag.delete_document(kb_type="client", kb_id=str(brand_id), filename=safe_path)
    if not result.get("success") and safe_path != _safe_filename(filename):
        result = rag.delete_document(kb_type="client", kb_id=str(brand_id), filename=_safe_filename(filename))

    if not result.get("success"):
        raise HTTPException(status_code=404, detail=result.get("error", "删除失败"))

    _logger.info(f"[client-kb] 删除文档: brand_id={brand_id}, filename={safe_path}")
    return {"success": True, "message": "文档已删除"}


# ==================== 公共知识库 API ====================

@router.get("/public", summary="获取公共知识库文档列表")
async def list_public_documents():
    """获取公共知识库文档列表"""
    rag = get_unified_rag()
    result = rag.list_documents(kb_type="public")

    if not result.get("success"):
        raise HTTPException(status_code=500, detail=result.get("error", "获取失败"))

    return {
        "success": True,
        "documents": result.get("documents", []),
        "total": result.get("total", 0)
    }


@router.post("/public/upload", summary="上传文档到公共知识库")
async def upload_public_document(request: Request, file: UploadFile = File(...)):
    """上传文档到公共知识库（支持 md/txt/json/pdf）
    [P1-15 fix 2026-05-23 老板授权] Codex 跨 AI 审计 · 加 is_admin 校验防普通代理污染全局语料"""
    user = getattr(request.state, "user", None)
    if not user or not user.get("is_admin"):
        raise HTTPException(status_code=403, detail="需要管理员权限")
    filename = file.filename or "unknown.txt"

    try:
        content = await file.read()

        if filename.lower().endswith('.pdf'):
            try:
                from services import pdf_reader
                content_str = "\n\n".join(pdf_reader.page_texts(content))
            except ImportError:
                raise HTTPException(status_code=500, detail="服务器未安装 PDF 解析库")
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"PDF 解析失败: {e}")
        else:
            content_str = content.decode("utf-8")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件读取失败: {e}")

    rag = get_unified_rag()
    result = rag.add_document(kb_type="public", kb_id="", filename=filename, content=content_str)

    if not result.get("success"):
        raise HTTPException(status_code=500, detail=result.get("error", "上传失败"))

    _logger.info(f"[public-kb] 上传文档: {filename}, size={result.get('size', 0)}")
    return {
        "success": True,
        "message": f"文档 {filename} 已上传到公共知识库",
        "path": result.get("path", ""),
        "size": result.get("size", 0)
    }


@router.delete("/public/{filename}", summary="删除公共知识库文档")
async def delete_public_document(filename: str, request: Request):
    """删除公共知识库文档
    [P1-15 fix 2026-05-23 老板授权] 加 is_admin 校验"""
    user = getattr(request.state, "user", None)
    if not user or not user.get("is_admin"):
        raise HTTPException(status_code=403, detail="需要管理员权限")
    rag = get_unified_rag()
    result = rag.delete_document(kb_type="public", filename=filename)

    if not result.get("success"):
        raise HTTPException(status_code=404, detail=result.get("error", "删除失败"))

    _logger.info(f"[public-kb] 删除文档: {filename}")
    return {"success": True, "message": "文档已删除"}


# ==================== 兼容前端通用知识库 API ====================

@router.get("/stats", summary="知识库统计")
async def get_knowledge_stats(request: Request):
    # [GEO-R1-CAN-064] get_stats 枚举全部客户目录(brand_id 为租户标识)+ 角色目录,
    # 属跨租户元数据枚举 · 限管理员,防低权用户拉全平台品牌清单与知识量。
    _require_admin(request)
    rag = get_unified_rag()
    return {"status": "success", "success": True, "stats": rag.get_stats()}


@router.get("/{kb_type}", summary="通用知识库文档列表")
async def list_knowledge(kb_type: str, request: Request, kb_id: str = "", role_type: str = ""):
    """兼容前端 knowledgeApi.list('/api/knowledge/{kb_type}?kb_id=...')。"""
    brand_id = None
    if kb_type == "client":
        if not kb_id:
            raise HTTPException(status_code=400, detail="client 知识库需要 kb_id")
        brand_id = _parse_client_brand_id(kb_id)
        from auth.brand_access import require_brand_access
        require_brand_access(request, brand_id)
    else:
        # [GEO-R1-CAN-057] role/public 为全局知识库命名空间,通用列举需管理员
        _require_admin(request)

    rag = get_unified_rag()
    result = rag.list_documents(kb_type=kb_type, kb_id=kb_id or None, role_type=role_type or None)
    if not result.get("success"):
        raise HTTPException(status_code=500, detail=result.get("error", "获取失败"))
    documents = result.get("documents", [])
    if kb_type == "client" and brand_id:
        documents = _enrich_client_documents(brand_id, documents)
    return {
        "status": "success",
        "success": True,
        "documents": documents,
        "total": result.get("total", 0),
    }


@router.delete("/{kb_type}/{kb_id}/{filename}", summary="通用知识库删除")
async def delete_knowledge(kb_type: str, kb_id: str, filename: str, request: Request, role_type: str = ""):
    if kb_type == "client":
        from auth.brand_access import require_brand_access
        require_brand_access(request, int(kb_id))
    else:
        # [Deploy-CTO NO-GO finding 3 返工 2026-07-12] 非 client 知识库(role/advisor/employee)是全站
        # 共享检索源,删除破坏性且跨租户 → 仅 admin。原实现仅对 client 校验,其余 kb_type 任意登录用户可删。
        _require_admin(request)

    rag = get_unified_rag()
    result = rag.delete_document(kb_type=kb_type, kb_id=kb_id, role_type=role_type or None, filename=_safe_filename(filename))
    if not result.get("success"):
        raise HTTPException(status_code=404, detail=result.get("error", "删除失败"))
    return {"status": "success", "success": True, "message": "文档已删除"}


@router.post("/search", summary="客户知识库语义检索")
async def search_knowledge(data: KnowledgeSearch, request: Request):
    if not data.query.strip():
        raise HTTPException(status_code=400, detail="查询内容为空")
    if not data.brand_id:
        raise HTTPException(status_code=400, detail="缺少 brand_id")

    from auth.brand_access import require_brand_access
    require_brand_access(request, data.brand_id)

    rag = get_unified_rag()
    results = await rag.retrieve(
        query=data.query,
        brand_id=data.brand_id,
        top_k=max(1, min(data.top_k or 5, 20)),
        use_client=True,
        use_role=False,
    )
    return {
        "status": "success",
        "success": True,
        "results": [
            {
                "content": r.get("content", ""),
                "filename": r.get("source") or r.get("filename") or "",
                "brand_id": data.brand_id,
                "keywords": r.get("keywords", ""),
                "score": r.get("score", 0),
            }
            for r in results
        ],
        "reranked": False,
    }
