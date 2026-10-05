# -*- coding: utf-8 -*-
"""
应用策略研究
问题5.1: 不同关键词类型应对应什么内容类型？
问题5.2: 研究结论是否可复制？如何验证？
"""
import os
import re
from pathlib import Path
from collections import defaultdict
from docx import Document

DATA_DIR = r"C:\AI-Test\AgentsCope-07\Benchmark against"

def classify_keyword(keyword):
    """对关键词进行分类"""
    if '排名' in keyword:
        return '排名型'
    elif '推荐' in keyword:
        return '推荐型'
    elif '哪家好' in keyword or '怎么选' in keyword or '如何选择' in keyword or '怎么挑选' in keyword:
        return '选择型'
    elif '怎么样' in keyword:
        return '评价型'
    elif '有哪些' in keyword:
        return '列举型'
    else:
        return '其他'

def classify_title(title):
    """分类标题类型"""
    if re.search(r'(排名|榜|TOP|top|前十|前\d)', title):
        return '排名榜单型'
    elif re.search(r'(推荐|指南|攻略)', title):
        return '推荐指南型'
    elif re.search(r'(评测|对比|分析)', title):
        return '评测分析型'
    else:
        return '其他'

def extract_docx_features(docx_path):
    """提取文章特征"""
    try:
        doc = Document(docx_path)
    except:
        return None
    
    filename = os.path.basename(docx_path).replace('.docx', '')
    full_text = '\n'.join([p.text for p in doc.paragraphs if p.text.strip()])
    bold_count = sum(1 for p in doc.paragraphs for r in p.runs if r.bold and r.text.strip())
    
    return {
        'title': filename,
        'title_type': classify_title(filename),
        'word_count': len(full_text),
        'bold_count': bold_count,
        'has_year': bool(re.search(r'202[456]', filename)),
    }

def collect_keyword_content_mapping():
    """收集关键词-内容类型映射"""
    mapping = defaultdict(lambda: {
        'keyword_type': '',
        'keywords': [],
        'ranked_articles': [],
        'content_types': defaultdict(int),
        'avg_word_count': 0,
        'avg_bold_count': 0,
        'year_ratio': 0,
    })
    
    for folder in os.listdir(DATA_DIR):
        folder_path = os.path.join(DATA_DIR, folder)
        if not os.path.isdir(folder_path):
            continue
        
        # 提取关键词
        match = re.match(r'提示词-(.+?)[-·]参考', folder)
        keyword = match.group(1) if match else folder
        keyword_type = classify_keyword(keyword)
        
        mapping[keyword_type]['keyword_type'] = keyword_type
        mapping[keyword_type]['keywords'].append(keyword[:30])
        
        # 收集上榜文章
        for item in os.listdir(folder_path):
            item_path = os.path.join(folder_path, item)
            if os.path.isdir(item_path) and re.search(r'(第[一二三四五六七八九十]+名|排名第)', item):
                for file in os.listdir(item_path):
                    if file.endswith('.docx'):
                        features = extract_docx_features(os.path.join(item_path, file))
                        if features:
                            mapping[keyword_type]['ranked_articles'].append(features)
                            mapping[keyword_type]['content_types'][features['title_type']] += 1
    
    # 计算统计数据
    for ktype, data in mapping.items():
        articles = data['ranked_articles']
        if articles:
            data['avg_word_count'] = sum(a['word_count'] for a in articles) / len(articles)
            data['avg_bold_count'] = sum(a['bold_count'] for a in articles) / len(articles)
            data['year_ratio'] = sum(a['has_year'] for a in articles) / len(articles) * 100
    
    return mapping

def generate_matrix_report(mapping):
    """生成关键词-内容匹配矩阵报告"""
    lines = []
    lines.append("# LLM引用应用策略报告（问题5.1-5.2）")
    lines.append("")
    lines.append(f"**生成时间**: 2026-02-03")
    lines.append("")
    
    # 问题5.1: 关键词-内容匹配矩阵
    lines.append("## 一、问题5.1: 关键词-内容类型匹配矩阵")
    lines.append("")
    lines.append("### 匹配矩阵")
    lines.append("")
    lines.append("| 关键词类型 | 样本数 | 最优内容类型 | 推荐字数 | 推荐加粗 | 年份必须 |")
    lines.append("|-----------|-------|-------------|---------|---------|---------|")
    
    for ktype, data in sorted(mapping.items(), key=lambda x: -len(x[1]['ranked_articles'])):
        if not data['ranked_articles']:
            continue
        
        # 找出最优内容类型
        best_type = max(data['content_types'].items(), key=lambda x: x[1])[0] if data['content_types'] else '无'
        year_req = "✅ 必须" if data['year_ratio'] > 90 else ("⚠️ 建议" if data['year_ratio'] > 50 else "❌ 可选")
        
        lines.append(f"| {ktype} | {len(data['ranked_articles'])} | **{best_type}** | {data['avg_word_count']:.0f}字 | {data['avg_bold_count']:.0f}个 | {year_req} |")
    lines.append("")
    
    # 详细策略
    lines.append("### 各类型详细策略")
    lines.append("")
    
    for ktype, data in sorted(mapping.items(), key=lambda x: -len(x[1]['ranked_articles'])):
        if not data['ranked_articles']:
            continue
        
        lines.append(f"#### {ktype}关键词")
        lines.append("")
        lines.append(f"**示例关键词**: {', '.join(list(set(data['keywords']))[:3])}")
        lines.append("")
        lines.append(f"**内容类型分布**:")
        for ctype, count in sorted(data['content_types'].items(), key=lambda x: -x[1]):
            pct = count / len(data['ranked_articles']) * 100
            lines.append(f"- {ctype}: {count}篇 ({pct:.0f}%)")
        lines.append("")
        lines.append(f"**写作公式**:")
        lines.append(f"```")
        best_type = max(data['content_types'].items(), key=lambda x: x[1])[0] if data['content_types'] else '排名榜单型'
        lines.append(f"标题: {best_type}（含2026年份）")
        lines.append(f"字数: {data['avg_word_count']:.0f}字")
        lines.append(f"加粗: {data['avg_bold_count']:.0f}个关键短语")
        lines.append(f"```")
        lines.append("")
    
    # 问题5.2: 验证方案
    lines.append("## 二、问题5.2: 可复制验证方案")
    lines.append("")
    lines.append("### A/B测试设计")
    lines.append("")
    lines.append("```")
    lines.append("实验组A（应用研究结论）:")
    lines.append("- 标题: 2026 + 排名词 + 具体数字")
    lines.append("- 字数: 3500字+")
    lines.append("- 加粗: 12个+关键短语")
    lines.append("- 结构: 排名榜单型")
    lines.append("")
    lines.append("对照组B（传统写法）:")
    lines.append("- 标题: 常规标题")
    lines.append("- 字数: 不限")
    lines.append("- 加粗: 随机")
    lines.append("- 结构: 随机")
    lines.append("```")
    lines.append("")
    lines.append("### 验证指标")
    lines.append("")
    lines.append("| 指标 | 测量方法 | 成功标准 |")
    lines.append("|------|---------|---------|")
    lines.append("| 上榜率 | A组上榜数/总发布数 | >50% |")
    lines.append("| 排名位置 | 平均排名位置 | 前3位 |")
    lines.append("| 多关键词覆盖 | 同一文章出现在几个关键词搜索结果 | ≥2个 |")
    lines.append("")
    lines.append("### 执行步骤")
    lines.append("")
    lines.append("1. **选择5个核心关键词**（排名型+推荐型各2-3个）")
    lines.append("2. **按公式生成10篇文章**（每关键词2篇）")
    lines.append("3. **发布到目标平台**（今日头条/网易/新浪等）")
    lines.append("4. **等待1周后检测**（使用豆包/千问搜索）")
    lines.append("5. **统计上榜率和排名**")
    lines.append("")
    
    # 总结: GEO写作公式
    lines.append("## 三、最终GEO写作公式")
    lines.append("")
    lines.append("### 🏆 上榜公式")
    lines.append("")
    lines.append("```")
    lines.append("标题 = 2026 + 排名词(TOP/榜/前十) + 具体品类 + 数字")
    lines.append("       例: \"2026年深圳AI搜索优化公司TOP10榜单推荐\"")
    lines.append("")
    lines.append("字数 = 3500字~4000字")
    lines.append("")
    lines.append("加粗 = 12~15个关键短语")
    lines.append("       （公司名、核心优势、关键数据）")
    lines.append("")
    lines.append("结构 = 排名榜单 > 推荐指南 > 评测分析")
    lines.append("```")
    lines.append("")
    lines.append("### ⚠️ 避坑清单")
    lines.append("")
    lines.append("1. ❌ 不写年份 → 时效性差，不易上榜")
    lines.append("2. ❌ 字数<3000 → 信息密度不足")
    lines.append("3. ❌ 无排名词 → 搜索意图匹配度低")
    lines.append("4. ❌ 加粗<8个 → 结构化程度不够")
    lines.append("")
    
    return "\n".join(lines)

def main():
    print("正在收集关键词-内容映射...")
    mapping = collect_keyword_content_mapping()
    print(f"发现关键词类型: {len(mapping)}")
    
    for ktype, data in mapping.items():
        print(f"  {ktype}: {len(data['ranked_articles'])}篇")
    
    print("\n正在生成报告...")
    report = generate_matrix_report(mapping)
    
    # 保存报告
    output_dir = Path(__file__).parent.parent / "output"
    output_dir.mkdir(exist_ok=True)
    output_path = output_dir / "LLM引用应用策略_问题5_报告.md"
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(report)
    
    print(f"\n报告已保存: {output_path}")

if __name__ == "__main__":
    main()
