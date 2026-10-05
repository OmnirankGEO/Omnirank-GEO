"""
诊断报告 MD -> PPTX 生成器
支持 GEO专项 / 社媒专项 / 全面诊断 三种scope
使用 python-pptx 生成专业 PPT 报告
"""

import re
import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Cm, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE


# ====== 配色方案 ======
BRAND_DARK = RGBColor(0x0F, 0x17, 0x2A)     # 深蓝背景
BRAND_BLUE = RGBColor(0x1E, 0x40, 0xAF)     # 主蓝
BRAND_LIGHT_BLUE = RGBColor(0x3B, 0x82, 0xF6)  # 亮蓝
BRAND_CYAN = RGBColor(0x06, 0xB6, 0xD4)     # 青色
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF1, 0xF5, 0xF9)
DARK_TEXT = RGBColor(0x1E, 0x29, 0x3B)
MED_TEXT = RGBColor(0x47, 0x55, 0x69)
RED = RGBColor(0xEF, 0x44, 0x44)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
GREEN = RGBColor(0x22, 0xC5, 0x5E)

SLIDE_WIDTH = Cm(33.867)  # 16:9
SLIDE_HEIGHT = Cm(19.05)
MARGIN_L = Cm(2.5)
MARGIN_R = Cm(2.5)
CONTENT_W = SLIDE_WIDTH - MARGIN_L - MARGIN_R


def add_bg(slide, color):
    """给 slide 添加纯色背景"""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text,
                font_size=12, bold=False, color=DARK_TEXT,
                alignment=PP_ALIGN.LEFT, font_name="Microsoft YaHei"):
    """添加文本框的快捷函数"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_paragraph(text_frame, text, font_size=12, bold=False,
                  color=DARK_TEXT, alignment=PP_ALIGN.LEFT, space_before=0):
    """向已有 text_frame 追加段落"""
    p = text_frame.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Microsoft YaHei"
    p.alignment = alignment
    if space_before:
        p.space_before = Pt(space_before)
    return p


def _scope_label(scope):
    """scope -> 中文报告标签"""
    return {'geo': 'GEO诊断报告', 'social': '社媒诊断报告', 'full': '全面诊断报告'}.get(scope, '诊断报告')


def add_footer(slide, brand_name, scope='geo'):
    """每页底部加 footer"""
    footer_h = Cm(1)
    footer_top = SLIDE_HEIGHT - footer_h
    # 左
    add_textbox(slide, MARGIN_L, footer_top, Cm(12), footer_h,
                "OmniRank AI | 全域上榜", font_size=7, color=MED_TEXT)
    # 右
    add_textbox(slide, SLIDE_WIDTH - MARGIN_R - Cm(12), footer_top, Cm(12), footer_h,
                f"{brand_name} {_scope_label(scope)}", font_size=7, color=MED_TEXT,
                alignment=PP_ALIGN.RIGHT)


def add_section_divider(prs, title, subtitle="", brand_name="", scope='geo'):
    """添加章节分割页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    add_bg(slide, BRAND_DARK)

    # 装饰圆
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL,
        SLIDE_WIDTH - Cm(8), Cm(-3), Cm(16), Cm(16))
    circle.fill.solid()
    circle.fill.fore_color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    circle.fill.fore_color.brightness = 0.0
    circle.line.fill.background()

    add_textbox(slide, MARGIN_L, Cm(6), CONTENT_W, Cm(3),
                title, font_size=36, bold=True, color=WHITE)
    if subtitle:
        add_textbox(slide, MARGIN_L, Cm(10), CONTENT_W, Cm(2),
                    subtitle, font_size=16, color=RGBColor(0x94, 0xA3, 0xB8))
    add_footer(slide, brand_name, scope)
    return slide


def add_table_slide(prs, title, headers, rows, brand_name="", scope='geo'):
    """添加带表格的 slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)

    # 标题
    add_textbox(slide, MARGIN_L, Cm(1.2), CONTENT_W, Cm(1.5),
                title, font_size=22, bold=True, color=BRAND_BLUE)

    # 表格
    cols = len(headers)
    tbl_rows = len(rows) + 1
    table_top = Cm(3.5)
    table_h = min(Cm(0.9) * tbl_rows, Cm(14))
    table = slide.shapes.add_table(tbl_rows, cols,
        MARGIN_L, table_top, CONTENT_W, table_h).table

    # 表头样式
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = BRAND_DARK
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.font.name = "Microsoft YaHei"

    # 数据行
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = str(val)
            if r_idx % 2 == 1:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_GRAY
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9)
                p.font.color.rgb = DARK_TEXT
                p.font.name = "Microsoft YaHei"

    add_footer(slide, brand_name, scope)
    return slide


# ====== 从 MD 中提取数据 ======
def detect_scope(md):
    """从报告内容检测scope类型"""
    if 'AI搜索可见度诊断报告' in md or 'GEO专项' in md:
        return 'geo'
    if '社媒内容生态诊断报告' in md or '社媒专项' in md:
        return 'social'
    return 'full'


def _strip_md(s):
    """Strip markdown bold markers and whitespace from extracted value"""
    return s.strip().strip('*').strip()


def extract_meta(md):
    scope = detect_scope(md)
    m_brand = re.search(r'品牌名[：:](.+)', md)
    m_industry = re.search(r'行业[：:](.+)', md)
    m_score = re.search(r'(?:GEO|社媒)?总分[：:].*?(\d+)', md)
    m_level = re.search(r'等级[：:](.+)', md)
    m_date = re.search(r'诊断日期[：:](.+)', md)
    m_ai_rate = re.search(r'AI推荐率[：:].*?([\d.]+)%', md)
    m_ai_count = re.search(r'AI测试次数[：:].*?(\d+)', md)
    m_ai_rec = re.search(r'品牌被推荐次数[：:].*?(\d+)', md)
    return {
        'brand': _strip_md(m_brand.group(1)) if m_brand else '品牌',
        'industry': _strip_md(m_industry.group(1)) if m_industry else '',
        'score': int(m_score.group(1)) if m_score else 0,
        'level': _strip_md(m_level.group(1)) if m_level else '',
        'date': _strip_md(m_date.group(1)) if m_date else '',
        'ai_rate': float(m_ai_rate.group(1)) if m_ai_rate else 0,
        'ai_count': int(m_ai_count.group(1)) if m_ai_count else 0,
        'ai_rec': int(m_ai_rec.group(1)) if m_ai_rec else 0,
        'scope': scope,
    }


def extract_score_table(md):
    """提取七维度评分表"""
    pattern = r'\| ([\U0001f000-\U0001ffff\w\s]+?) \| (\d+) \| (\d+) \| (.+?) \|'
    matches = re.findall(pattern, md)
    scores = []
    for name, score, total, status in matches:
        name = name.strip()
        if any(k in name for k in [
            # GEO scope
            'AI引擎', 'AI推荐', '网页', '权威', '结构化', '品牌基础',
            # Social scope
            '品牌存在感', '竞品活跃', '行业生态', '内容质量', '运营基础',
            # Full scope
            '社媒', '品牌',
        ]):
            scores.append({
                'name': name,
                'score': int(score),
                'total': int(total),
                'status': status.strip(),
            })
    return scores


def extract_competitors(md):
    """提取竞品表格"""
    # 社媒竞品
    pattern = r'\| (.+?) \| (抖音|小红书) \| (.+?) \| (.+?) \| (.+?) \|'
    matches = re.findall(pattern, md)
    comps = []
    for name, platform, fans, style, advantage in matches:
        name = name.strip()
        if name and '账号名称' not in name and '---' not in name:
            comps.append([name, platform, fans.strip(), advantage.strip()])
    return comps


def extract_action_plan(md):
    """提取30天速赢计划"""
    pattern = r'\| (\d+) \| (.+?) \| (.+?) \| (.+?) \|'
    matches = re.findall(pattern, md)
    actions = []
    for num, action, person, effect in matches:
        action = action.strip()
        if action and '具体行动' not in action:
            actions.append([num, action[:60], effect.strip()[:50]])
    return actions[:10]


# ====== 从 LLM JSON 数据转换为内部格式 ======
def _convert_json_data(json_data):
    """将 extract_report_data.py 输出的 JSON 转为 generate_pptx 内部格式"""
    ai_tests = json_data.get('aiTests', 0)
    ai_rate = json_data.get('aiRecommendRate', 0)
    meta = {
        'brand': json_data.get('brandName', '品牌'),
        'industry': json_data.get('industry', ''),
        'score': json_data.get('geoScore', 0),
        'level': json_data.get('level', ''),
        'date': json_data.get('date', ''),
        'ai_rate': ai_rate,
        'ai_count': ai_tests,
        'ai_rec': round(ai_tests * ai_rate / 100) if ai_tests else 0,
        'scope': json_data.get('scope', 'geo'),
        'content_total': json_data.get('contentTotal', 0),
        'platform_count': json_data.get('platformCount', 0),
    }
    scores = [
        {'name': d.get('name', ''), 'score': d.get('score', 0),
         'total': d.get('max', 0), 'status': d.get('status', '')}
        for d in json_data.get('dims', [])
    ]
    comps = [
        [c.get('name', ''), c.get('tag', ''), c.get('strength', ''),
         (c.get('points', []) or [''])[0]]
        for c in json_data.get('competitors', [])
    ]
    actions_raw = []
    idx = 1
    for phase in json_data.get('actions', []):
        for item in phase.get('items', []):
            actions_raw.append([str(idx), item.get('task', ''), item.get('detail', '')])
            idx += 1
    findings = [
        (f.get('status', 'warning'), f.get('title', ''), f.get('desc', ''))
        for f in json_data.get('findings', [])
    ]
    quick_wins = json_data.get('executiveSummary', {}).get('quickWins', [])
    recs = [
        (str(i + 1), qw.get('action', ''), '', qw.get('impact', ''))
        for i, qw in enumerate(quick_wins)
    ]
    return meta, scores, comps, actions_raw, findings, recs


# ====== 主流程 ======
def generate_pptx(md_path, output_path, json_data=None):
    print(f"[PPTX] Reading report: {md_path}")
    md = open(md_path, 'r', encoding='utf-8').read()

    if json_data:
        print(f"[PPTX] Using LLM-extracted JSON data ({json_data.get('brandName', '?')})")
        meta, scores, comps, actions, json_findings, json_recs = _convert_json_data(json_data)
    else:
        meta = extract_meta(md)
        scores = extract_score_table(md)
        comps = extract_competitors(md)
        actions = extract_action_plan(md)
        json_findings = None
        json_recs = None

    scope = meta.get('scope', 'geo')
    # Enrich meta with social-specific metrics from MD
    if scope == 'social':
        m_ct = re.search(r'品牌内容总量\s*\|\s*(\d+)', md)
        m_pc = re.search(r'平台覆盖\s*\|\s*(\d+)', md)
        meta['content_total'] = int(m_ct.group(1)) if m_ct else 0
        meta['platform_count'] = int(m_pc.group(1)) if m_pc else 0
    print(f"[PPTX] Brand: {meta['brand']} | Score: {meta['score']} | Level: {meta['level']} | Scope: {scope}")

    # scope-aware labels
    _labels = {
        'geo':    {'tag': 'GEO 诊断报告', 'ver': 'GEO诊断系统 v5.0',
                   'overview': '向4大主流AI平台提出真实用户场景问题，综合评估品牌AI可见度',
                   'score': 'GEO综合评分', 'section': 'GEO评分详解',
                   'sub': '基于全域上榜v5.0评分体系，满分100分',
                   'detail': '评分详情', 'table': 'GEO评分维度明细',
                   'cta': '开始提升您的AI可见度',
                   'tagline': '让您的品牌在AI搜索时代被看见'},
        'social': {'tag': '社媒 诊断报告', 'ver': '社媒诊断系统 v5.0',
                   'overview': '全面评估品牌在抖音、小红书等主流社交媒体平台的内容生态与竞争力',
                   'score': '社媒综合评分', 'section': '社媒评分详解',
                   'sub': '基于全域上榜v5.0评分体系，满分100分',
                   'detail': '评分详情', 'table': '社媒评分维度明细',
                   'cta': '开始提升您的社媒影响力',
                   'tagline': '让您的品牌在社交媒体时代脱颖而出'},
        'full':   {'tag': '全面 诊断报告', 'ver': '全域诊断系统 v5.0',
                   'overview': '综合评估品牌AI搜索可见度与社媒内容生态，全面把脉数字化竞争力',
                   'score': '综合评分', 'section': '评分详解',
                   'sub': '基于全域上榜v5.0七维度评分体系，满分100分',
                   'detail': '七维度评分详情', 'table': '评分维度明细',
                   'cta': '开始提升您的AI可见度与社媒影响力',
                   'tagline': '让您的品牌在AI搜索与社交媒体时代被看见'},
    }
    L = _labels.get(scope, _labels['geo'])
    dim_count = len(scores) if scores else 5

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # ========== SLIDE 1: 封面 ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BRAND_DARK)

    # 装饰
    circle1 = slide.shapes.add_shape(MSO_SHAPE.OVAL,
        SLIDE_WIDTH - Cm(10), Cm(-5), Cm(20), Cm(20))
    circle1.fill.solid()
    circle1.fill.fore_color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    circle1.line.fill.background()

    circle2 = slide.shapes.add_shape(MSO_SHAPE.OVAL,
        Cm(-5), SLIDE_HEIGHT - Cm(8), Cm(14), Cm(14))
    circle2.fill.solid()
    circle2.fill.fore_color.rgb = RGBColor(0x15, 0x25, 0x45)
    circle2.line.fill.background()

    # 诊断报告标签
    add_textbox(slide, MARGIN_L, Cm(3), Cm(20), Cm(1),
                L['tag'], font_size=14, color=RGBColor(0x94, 0xA3, 0xB8))

    # 品牌名
    add_textbox(slide, MARGIN_L, Cm(4.5), Cm(25), Cm(3),
                meta['brand'], font_size=48, bold=True, color=WHITE)

    # 行业
    add_textbox(slide, MARGIN_L, Cm(8.5), Cm(25), Cm(1.5),
                meta['industry'], font_size=16, color=RGBColor(0x94, 0xA3, 0xB8))

    # 分数框
    score_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
        MARGIN_L, Cm(11), Cm(14), Cm(4))
    score_box.fill.solid()
    score_box.fill.fore_color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    score_box.line.color.rgb = RGBColor(0x33, 0x44, 0x55)
    score_box.line.width = Pt(1)

    # 分数数字
    add_textbox(slide, Cm(3.5), Cm(11.3), Cm(5), Cm(3.5),
                str(meta['score']), font_size=64, bold=True, color=BRAND_LIGHT_BLUE)
    add_textbox(slide, Cm(8), Cm(12.5), Cm(3), Cm(1.5),
                "/100", font_size=20, color=RGBColor(0x64, 0x74, 0x8B))

    # 等级标签
    level_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
        Cm(11.5), Cm(12.2), Cm(4), Cm(1.5))
    level_box.fill.solid()
    level_box.fill.fore_color.rgb = RGBColor(0x2A, 0x1F, 0x05)
    level_box.line.color.rgb = AMBER
    level_box.line.width = Pt(1)
    add_textbox(slide, Cm(11.5), Cm(12.2), Cm(4), Cm(1.5),
                meta['level'], font_size=14, bold=True,
                color=AMBER, alignment=PP_ALIGN.CENTER)

    # 底部信息
    add_textbox(slide, MARGIN_L, Cm(16.5), Cm(12), Cm(1),
                "OmniRank AI | 全域上榜", font_size=14, bold=True,
                color=RGBColor(0x64, 0x74, 0x8B))
    add_textbox(slide, MARGIN_L, Cm(17.5), Cm(12), Cm(0.8),
                f"诊断日期：{meta['date']}  |  {L['ver']}",
                font_size=10, color=RGBColor(0x64, 0x74, 0x8B))
    add_textbox(slide, SLIDE_WIDTH - MARGIN_R - Cm(8), Cm(17.5), Cm(8), Cm(0.8),
                "www.omnirank.cn", font_size=10,
                color=RGBColor(0x64, 0x74, 0x8B), alignment=PP_ALIGN.RIGHT)

    # ========== SLIDE 2: 诊断概览 ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    add_textbox(slide, MARGIN_L, Cm(1.2), CONTENT_W, Cm(1.5),
                "诊断概览", font_size=26, bold=True, color=BRAND_BLUE)
    add_textbox(slide, MARGIN_L, Cm(2.8), CONTENT_W, Cm(1),
                L['overview'],
                font_size=12, color=MED_TEXT)

    # 三个指标卡片 — scope-specific
    card_w = Cm(9)
    card_h = Cm(5)
    card_top = Cm(5)
    if scope == 'social':
        cards_data = [
            ("品牌内容总量", f"{meta.get('content_total', 0)}条", BRAND_LIGHT_BLUE),
            ("平台覆盖", f"{meta.get('platform_count', 0)}个平台", AMBER),
            (L['score'], f"{meta['score']}/100", RED if meta['score'] < 40 else AMBER),
        ]
    else:
        cards_data = [
            ("AI推荐次数", f"{meta['ai_rec']}次/{meta['ai_count']}次", BRAND_LIGHT_BLUE),
            ("AI推荐率", f"{meta['ai_rate']}%", AMBER),
            (L['score'], f"{meta['score']}/100", RED if meta['score'] < 40 else AMBER),
        ]
    for i, (label, value, accent) in enumerate(cards_data):
        left = MARGIN_L + Cm(i * 10)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
            left, card_top, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_GRAY
        card.line.fill.background()
        add_textbox(slide, left + Cm(1), card_top + Cm(0.8), Cm(7), Cm(1),
                    label, font_size=11, color=MED_TEXT)
        add_textbox(slide, left + Cm(1), card_top + Cm(2), Cm(7), Cm(2),
                    value, font_size=28, bold=True, color=accent)

    # 核心评分维度 — 使用实际评分数据
    dim_section_title = "三维度AI测试结果" if scope == 'geo' else "核心评分维度"
    add_textbox(slide, MARGIN_L, Cm(11.5), CONTENT_W, Cm(1),
                dim_section_title, font_size=16, bold=True, color=DARK_TEXT)

    if scores and len(scores) >= 3:
        for i, s in enumerate(scores[:3]):
            name = re.sub(r'[\U0001f000-\U0001ffff]\s*', '', s['name'])
            left = MARGIN_L + Cm(i * 10)
            pct = s['score'] / s['total'] if s['total'] else 0
            sc = GREEN if pct >= 0.6 else AMBER if pct >= 0.3 else RED
            add_textbox(slide, left, Cm(13), Cm(9), Cm(0.8),
                        name, font_size=13, bold=True, color=DARK_TEXT)
            add_textbox(slide, left, Cm(13.8), Cm(9), Cm(0.6),
                        s.get('status', '')[:30], font_size=9, color=MED_TEXT)
            add_textbox(slide, left, Cm(14.5), Cm(9), Cm(1.2),
                        f"{s['score']}/{s['total']}", font_size=24, bold=True, color=sc)
    else:
        fallback_dims = [
            ("品牌存在感", "品牌在社媒的内容覆盖"),
            ("竞品活跃度", "竞品社媒运营强度"),
            ("内容质量", "爆款内容占比"),
        ] if scope == 'social' else [
            ("AI推荐率", "AI引擎推荐品牌频次"),
            ("网页内容", "品牌网页资产覆盖"),
            ("品牌基础", "品牌基础信息完整度"),
        ]
        for i, (dim, desc) in enumerate(fallback_dims):
            left = MARGIN_L + Cm(i * 10)
            add_textbox(slide, left, Cm(13), Cm(9), Cm(0.8),
                        dim, font_size=13, bold=True, color=DARK_TEXT)
            add_textbox(slide, left, Cm(13.8), Cm(9), Cm(0.6),
                        desc, font_size=9, color=MED_TEXT)
            add_textbox(slide, left, Cm(14.5), Cm(9), Cm(1.2),
                        "—", font_size=24, bold=True, color=MED_TEXT)

    add_footer(slide, meta['brand'], scope)

    # ========== SLIDE 3: 评分详解 (柱状图) ==========
    if scores:
        add_section_divider(prs, L['section'],
                            L['sub'],
                            meta['brand'], scope)

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_bg(slide, WHITE)
        add_textbox(slide, MARGIN_L, Cm(1.2), CONTENT_W, Cm(1.5),
                    L['detail'], font_size=22, bold=True, color=BRAND_BLUE)

        # 柱状图
        chart_data = CategoryChartData()
        chart_data.categories = [s['name'].replace('🤖 ', '').replace('📱 ', '').replace('🌐 ', '')
                                  .replace('🏛️ ', '').replace('💡 ', '').replace('📝 ', '')
                                  .replace('🏷️ ', '') for s in scores]
        chart_data.add_series('得分', [s['score'] for s in scores])
        chart_data.add_series('满分', [s['total'] - s['score'] for s in scores])

        chart_frame = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_STACKED, MARGIN_L, Cm(3.5),
            CONTENT_W, Cm(10), chart_data)
        chart = chart_frame.chart
        chart.has_legend = True

        # 图表颜色
        series_score = chart.series[0]
        series_score.format.fill.solid()
        series_score.format.fill.fore_color.rgb = BRAND_LIGHT_BLUE

        series_total = chart.series[1]
        series_total.format.fill.solid()
        series_total.format.fill.fore_color.rgb = RGBColor(0xE2, 0xE8, 0xF0)

        # 总分文字
        add_textbox(slide, MARGIN_L, Cm(14.5), CONTENT_W, Cm(1.5),
                    f"总分：{meta['score']}/100  |  等级：{meta['level']}",
                    font_size=16, bold=True, color=BRAND_BLUE, alignment=PP_ALIGN.CENTER)
        add_footer(slide, meta['brand'], scope)

    # ========== SLIDE 4: 评分明细表 ==========
    if scores:
        add_table_slide(prs, L['table'],
            ["维度", "得分", "满分", "评价"],
            [[s['name'], str(s['score']), str(s['total']), s['status']] for s in scores],
            meta['brand'], scope)

    # ========== SLIDE 5: 竞品分析 ==========
    if comps:
        if scope == 'geo':
            comp_subtitle = "AI推荐品牌对标分析"
            comp_title = "AI推荐竞品"
        elif scope == 'social':
            comp_subtitle = "社媒热门账号对标分析"
            comp_title = "社媒热门竞品"
        else:
            comp_subtitle = "AI推荐品牌 & 社媒热门账号对标分析"
            comp_title = "竞品对标"
        add_section_divider(prs, "竞品对标分析",
                            comp_subtitle, meta['brand'], scope)
        add_table_slide(prs, comp_title,
            ["账号名称", "平台", "粉丝量级", "核心优势"],
            comps[:7], meta['brand'], scope)

    # ========== SLIDE 6: 关键发现 ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    add_textbox(slide, MARGIN_L, Cm(1.2), CONTENT_W, Cm(1.5),
                "关键发现", font_size=22, bold=True, color=BRAND_BLUE)

    if json_findings:
        findings = json_findings
    else:
        findings_pattern = r'\| (🔴|🟡|🟢).+?\| \*\*(.+?)\*\* \| (.+?) \|'
        findings = re.findall(findings_pattern, md)
    for i, (icon_or_status, title, desc) in enumerate(findings[:3]):
        top = Cm(3.5) + Cm(i * 4)
        color_map = {'🔴': RED, '🟡': AMBER, '🟢': GREEN,
                     'danger': RED, 'warning': AMBER, 'good': GREEN}
        accent = color_map.get(icon_or_status, MED_TEXT)

        # 色条
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
            MARGIN_L, top, Cm(0.4), Cm(3))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

        add_textbox(slide, MARGIN_L + Cm(1), top + Cm(0.2), CONTENT_W - Cm(1), Cm(1),
                    title, font_size=14, bold=True, color=DARK_TEXT)
        add_textbox(slide, MARGIN_L + Cm(1), top + Cm(1.3), CONTENT_W - Cm(1), Cm(1.5),
                    desc.strip(), font_size=10, color=MED_TEXT)

    add_footer(slide, meta['brand'], scope)

    # ========== SLIDE 7: 优先建议 ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    add_textbox(slide, MARGIN_L, Cm(1.2), CONTENT_W, Cm(1.5),
                "优先建议（7天内可执行）", font_size=22, bold=True, color=BRAND_BLUE)

    if json_recs:
        recs = json_recs
    else:
        rec_pattern = r'\| (\d) \| \*\*(.+?)\*\*[：:](.+?) \| (.+?) \|'
        recs = re.findall(rec_pattern, md)
    for i, (num, title, detail, effect) in enumerate(recs[:3]):
        top = Cm(3.5) + Cm(i * 4.5)
        # 编号圆
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL,
            MARGIN_L, top, Cm(1.5), Cm(1.5))
        circle.fill.solid()
        circle.fill.fore_color.rgb = BRAND_LIGHT_BLUE
        circle.line.fill.background()
        add_textbox(slide, MARGIN_L, top, Cm(1.5), Cm(1.5),
                    num, font_size=18, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)

        add_textbox(slide, MARGIN_L + Cm(2), top + Cm(0.1), CONTENT_W - Cm(2), Cm(1),
                    title.strip(), font_size=13, bold=True, color=DARK_TEXT)
        add_textbox(slide, MARGIN_L + Cm(2), top + Cm(1.2), CONTENT_W - Cm(2), Cm(1),
                    detail.strip()[:80], font_size=9, color=MED_TEXT)
        add_textbox(slide, MARGIN_L + Cm(2), top + Cm(2.2), CONTENT_W - Cm(2), Cm(1),
                    f"预期效果：{effect.strip()[:60]}", font_size=9, bold=True, color=GREEN)

    add_footer(slide, meta['brand'], scope)

    # ========== SLIDE 8: 30天速赢计划 ==========
    if actions:
        add_section_divider(prs, "30天速赢计划",
                            "关键执行行动清单", meta['brand'], scope)
        add_table_slide(prs, "30天速赢计划",
            ["序号", "具体行动", "预期效果"],
            actions, meta['brand'], scope)

    # ========== SLIDE 9: 差距对标 ==========
    gap_pattern = r'\| (.+?) \| (.+?) \| (.+?) \| (.+?) \|\n'
    gaps = re.findall(gap_pattern, md[md.find('差距对标'):] if '差距对标' in md else '')
    gap_rows = []
    for dim, you, benchmark, gap in gaps:
        dim = dim.strip()
        if dim and '对比维度' not in dim and '---' not in dim:
            gap_rows.append([dim, you.strip(), benchmark.strip(), gap.strip()])
    if gap_rows:
        add_table_slide(prs, "差距对标",
            ["对比维度", "您的品牌", "行业标杆", "差距评估"],
            gap_rows, meta['brand'], scope)

    # ========== SLIDE 10: 联系我们 ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BRAND_DARK)

    add_textbox(slide, MARGIN_L, Cm(4), CONTENT_W, Cm(2),
                L['cta'], font_size=36, bold=True,
                color=WHITE, alignment=PP_ALIGN.CENTER)

    consultant = '社媒顾问' if scope == 'social' else 'GEO顾问'
    steps = [
        ("1", f"扫码添加专属{consultant}"),
        ("2", "获取报告一对一深度解读"),
        ("3", "定制专属优化执行方案"),
    ]
    for i, (num, text) in enumerate(steps):
        left = MARGIN_L + Cm(3) + Cm(i * 9)
        top = Cm(8)
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL,
            left + Cm(2.5), top, Cm(2), Cm(2))
        circle.fill.solid()
        circle.fill.fore_color.rgb = BRAND_LIGHT_BLUE
        circle.line.fill.background()
        add_textbox(slide, left + Cm(2.5), top, Cm(2), Cm(2),
                    num, font_size=20, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
        add_textbox(slide, left, top + Cm(2.5), Cm(7), Cm(1.5),
                    text, font_size=13, color=RGBColor(0xCB, 0xD5, 0xE1),
                    alignment=PP_ALIGN.CENTER)

    add_textbox(slide, MARGIN_L, Cm(14), CONTENT_W, Cm(1.5),
                "全域上榜（OmniRank AI）", font_size=18, bold=True,
                color=RGBColor(0x64, 0x74, 0x8B), alignment=PP_ALIGN.CENTER)
    add_textbox(slide, MARGIN_L, Cm(15.5), CONTENT_W, Cm(1),
                f"{L['tagline']}  |  www.omnirank.cn",
                font_size=12, color=RGBColor(0x64, 0x74, 0x8B), alignment=PP_ALIGN.CENTER)

    # ========== 保存 ==========
    prs.save(output_path)
    file_size = os.path.getsize(output_path) / 1024
    print(f"[PPTX] Generated: {output_path}")
    print(f"[PPTX] File size: {file_size:.1f} KB")
    print(f"[PPTX] Total slides: {len(prs.slides)}")


if __name__ == '__main__':
    md_path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(__file__), '..', 'output',
        '深圳美栖堂_upgrade_117_20260315001305.md')
    out_path = sys.argv[2] if len(sys.argv) > 2 else md_path.replace('.md', '.pptx')
    generate_pptx(md_path, out_path)
