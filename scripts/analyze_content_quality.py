# -*- coding: utf-8 -*-
"""
内容质量深度研究
分析正文内容的7个维度：
6.1 信息密度
6.2 权威引用
6.3 数据呈现
6.4 公司介绍深度
6.5 差异化标签
6.6 案例引用
6.7 评分体系
"""
import os
import re
from pathlib import Path
from collections import defaultdict
from docx import Document

DATA_DIR = r"C:\AI-Test\AgentsCope-07\Benchmark against"

# 权威来源关键词
AUTHORITY_SOURCES = [
    "普林斯顿", "Gartner", "艾瑞", "IDC", "Forrester", "麦肯锡", "贝恩",
    "中国信通院", "36氪", "亿欧", "波士顿咨询", "德勤", "毕马威", "安永",
    "哈佛", "斯坦福", "MIT", "清华", "北大", "中科院", "工信部", "信通院",
    "白皮书", "研究报告", "行业报告", "官方数据"
]

# 案例相关关键词
CASE_KEYWORDS = ["案例", "客户", "合作", "项目", "成功", "效果", "实践", "服务过"]

# 评分相关关键词
SCORE_KEYWORDS = ["评分", "分数", "满分", "得分", "打分", "评级", "星", "★"]

# 差异化标签关键词
DIFF_TAGS = [
    "第一", "领先", "头部", "专业", "技术派", "效果派", "服务派",
    "性价比", "全栈", "垂直", "深耕", "独家", "首创", "唯一"
]

def extract_content_features(docx_path):
    """提取正文内容的7个维度特征"""
    try:
        doc = Document(docx_path)
    except:
        return None
    
    filename = os.path.basename(docx_path).replace('.docx', '')
    
    # 收集所有段落文本
    paragraphs = []
    full_text = ""
    for para in doc.paragraphs:
        text = para.text.strip()
        if text and len(text) > 10:  # 过滤太短的段落
            paragraphs.append(text)
            full_text += text + "\n"
    
    if not full_text:
        return None
    
    # 6.1 信息密度
    para_count = len(paragraphs)
    avg_para_length = len(full_text) / para_count if para_count > 0 else 0
    
    # 6.2 权威引用
    authority_count = 0
    authority_sources_found = []
    for source in AUTHORITY_SOURCES:
        count = full_text.count(source)
        if count > 0:
            authority_count += count
            authority_sources_found.append(source)
    
    # 6.3 数据呈现
    # 统计表格数（通过docx表格）
    table_count = len(doc.tables)
    # 统计数字出现频率
    numbers = re.findall(r'\d+(?:\.\d+)?%?', full_text)
    number_count = len(numbers)
    
    # 6.4 公司介绍深度 - 识别公司段落
    company_sections = re.findall(r'(?:第[一二三四五六七八九十]+名|TOP\s*\d+|排名第\d+)[：:\s]*([^\n]{10,200})', full_text)
    company_intro_count = len(company_sections)
    company_intro_avg_len = sum(len(s) for s in company_sections) / len(company_sections) if company_sections else 0
    
    # 6.5 差异化标签
    diff_tags_found = []
    for tag in DIFF_TAGS:
        if tag in full_text:
            diff_tags_found.append(tag)
    
    # 6.6 案例引用
    case_mentions = 0
    for kw in CASE_KEYWORDS:
        case_mentions += full_text.count(kw)
    
    # 6.7 评分体系
    has_score_system = False
    score_mentions = 0
    for kw in SCORE_KEYWORDS:
        count = full_text.count(kw)
        if count > 0:
            score_mentions += count
            has_score_system = True
    
    return {
        'title': filename[:50],
        'word_count': len(full_text),
        
        # 6.1 信息密度
        'para_count': para_count,
        'avg_para_length': avg_para_length,
        
        # 6.2 权威引用
        'authority_count': authority_count,
        'authority_sources': authority_sources_found,
        
        # 6.3 数据呈现
        'table_count': table_count,
        'number_count': number_count,
        
        # 6.4 公司介绍深度
        'company_intro_count': company_intro_count,
        'company_intro_avg_len': company_intro_avg_len,
        
        # 6.5 差异化标签
        'diff_tags': diff_tags_found,
        'diff_tag_count': len(diff_tags_found),
        
        # 6.6 案例引用
        'case_mentions': case_mentions,
        
        # 6.7 评分体系
        'has_score_system': has_score_system,
        'score_mentions': score_mentions,
    }

def collect_all_articles():
    """收集所有文章"""
    ranked_articles = []
    cited_articles = []
    
    for folder in os.listdir(DATA_DIR):
        folder_path = os.path.join(DATA_DIR, folder)
        if not os.path.isdir(folder_path):
            continue
        
        for item in os.listdir(folder_path):
            item_path = os.path.join(folder_path, item)
            
            if item.endswith('.docx') and '回答' in item:
                continue
            
            if os.path.isdir(item_path):
                is_ranked = bool(re.search(r'(第[一二三四五六七八九十]+名|排名第)', item))
                
                for file in os.listdir(item_path):
                    if file.endswith('.docx'):
                        features = extract_content_features(os.path.join(item_path, file))
                        if features:
                            features['is_ranked'] = is_ranked
                            if is_ranked:
                                ranked_articles.append(features)
                            else:
                                cited_articles.append(features)
    
    return ranked_articles, cited_articles

def analyze_dimension(articles, dim_name, key):
    """分析单个维度的统计数据"""
    values = [a[key] for a in articles if a.get(key) is not None]
    if not values:
        return {'avg': 0, 'min': 0, 'max': 0, 'count': 0}
    
    if isinstance(values[0], bool):
        return {'true_ratio': sum(values) / len(values) * 100, 'count': len(values)}
    elif isinstance(values[0], (int, float)):
        return {
            'avg': sum(values) / len(values),
            'min': min(values),
            'max': max(values),
            'count': len(values)
        }
    elif isinstance(values[0], list):
        all_items = []
        for v in values:
            all_items.extend(v)
        from collections import Counter
        return {'top_items': Counter(all_items).most_common(10), 'count': len(values)}
    return {}

def generate_report(ranked, cited):
    """生成内容质量分析报告"""
    lines = []
    lines.append("# LLM引用内容质量深度研究报告（问题6.1-6.7）")
    lines.append("")
    lines.append(f"**生成时间**: 2026-02-03")
    lines.append(f"**样本量**: 上榜{len(ranked)}篇 vs 仅引用{len(cited)}篇")
    lines.append("")
    
    # 汇总对比表
    lines.append("## 一、上榜 vs 仅引用 核心差异一览")
    lines.append("")
    lines.append("| 维度 | 指标 | 上榜均值 | 仅引用均值 | 差异 | 结论 |")
    lines.append("|------|------|---------|-----------|------|------|")
    
    comparisons = [
        ('6.1 信息密度', 'para_count', '段落数'),
        ('6.1 信息密度', 'avg_para_length', '平均段落长度'),
        ('6.2 权威引用', 'authority_count', '权威引用次数'),
        ('6.3 数据呈现', 'table_count', '表格数'),
        ('6.3 数据呈现', 'number_count', '数字数据量'),
        ('6.4 公司介绍', 'company_intro_count', '公司介绍数'),
        ('6.5 差异化标签', 'diff_tag_count', '差异化标签数'),
        ('6.6 案例引用', 'case_mentions', '案例提及次数'),
        ('6.7 评分体系', 'score_mentions', '评分关键词'),
    ]
    
    for dim, key, label in comparisons:
        r_stats = analyze_dimension(ranked, dim, key)
        c_stats = analyze_dimension(cited, dim, key)
        
        r_avg = r_stats.get('avg', 0)
        c_avg = c_stats.get('avg', 0)
        diff = r_avg - c_avg
        
        if diff > 0:
            conclusion = "✅ 上榜更高"
        elif diff < 0:
            conclusion = "❌ 引用更高"
        else:
            conclusion = "➖ 相近"
        
        lines.append(f"| {dim} | {label} | {r_avg:.1f} | {c_avg:.1f} | {diff:+.1f} | {conclusion} |")
    
    lines.append("")
    
    # 6.2 权威引用详细分析
    lines.append("## 二、问题6.2: 权威引用分析")
    lines.append("")
    r_sources = analyze_dimension(ranked, '权威来源', 'authority_sources')
    lines.append("### 上榜文章最常引用的权威来源")
    lines.append("")
    if r_sources.get('top_items'):
        lines.append("| 排名 | 来源 | 出现篇数 |")
        lines.append("|------|------|---------|")
        for i, (source, count) in enumerate(r_sources['top_items'][:10], 1):
            lines.append(f"| {i} | {source} | {count} |")
    else:
        lines.append("*未发现明显的权威引用*")
    lines.append("")
    
    # 6.5 差异化标签分析
    lines.append("## 三、问题6.5: 差异化标签分析")
    lines.append("")
    r_tags = analyze_dimension(ranked, '差异化标签', 'diff_tags')
    lines.append("### 上榜文章最常用的差异化标签")
    lines.append("")
    if r_tags.get('top_items'):
        lines.append("| 排名 | 标签 | 出现篇数 |")
        lines.append("|------|------|---------|")
        for i, (tag, count) in enumerate(r_tags['top_items'][:10], 1):
            lines.append(f"| {i} | {tag} | {count} |")
    lines.append("")
    
    # 6.7 评分体系分析
    lines.append("## 四、问题6.7: 评分体系分析")
    lines.append("")
    r_score_ratio = sum(1 for a in ranked if a.get('has_score_system')) / len(ranked) * 100 if ranked else 0
    c_score_ratio = sum(1 for a in cited if a.get('has_score_system')) / len(cited) * 100 if cited else 0
    lines.append(f"- 上榜文章含评分体系比例: **{r_score_ratio:.1f}%**")
    lines.append(f"- 仅引用文章含评分体系比例: **{c_score_ratio:.1f}%**")
    lines.append("")
    
    # 关键洞察
    lines.append("## 五、关键洞察与内容公式")
    lines.append("")
    lines.append("### 🏆 上榜文章的内容特征公式")
    lines.append("")
    lines.append("```")
    
    r_para_avg = analyze_dimension(ranked, '', 'para_count').get('avg', 0)
    r_auth_avg = analyze_dimension(ranked, '', 'authority_count').get('avg', 0)
    r_table_avg = analyze_dimension(ranked, '', 'table_count').get('avg', 0)
    r_case_avg = analyze_dimension(ranked, '', 'case_mentions').get('avg', 0)
    
    lines.append(f"段落数: {r_para_avg:.0f}个（信息密度适中）")
    lines.append(f"权威引用: {r_auth_avg:.0f}次（增强可信度）")
    lines.append(f"表格数: {r_table_avg:.0f}个（结构化呈现）")
    lines.append(f"案例提及: {r_case_avg:.0f}次（增强说服力）")
    lines.append("```")
    lines.append("")
    
    return "\n".join(lines)

def main():
    print("正在收集文章...")
    ranked, cited = collect_all_articles()
    print(f"上榜文章: {len(ranked)}篇")
    print(f"仅引用文章: {len(cited)}篇")
    
    print("\n正在分析内容特征...")
    report = generate_report(ranked, cited)
    
    # 保存报告
    output_dir = Path(__file__).parent.parent / "output"
    output_dir.mkdir(exist_ok=True)
    output_path = output_dir / "LLM引用内容质量_问题6_报告.md"
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(report)
    
    print(f"\n报告已保存: {output_path}")

if __name__ == "__main__":
    main()
